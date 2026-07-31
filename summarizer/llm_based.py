"""LLM-backed summarizer using Gemini structured output and source labels."""
from __future__ import annotations

import json
import os
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable

from readers import extraction_is_sparse, read_file

from .base import Summary, Summarizer


OVERVIEW_HEADING = "기업 개요"
THESIS_HEADING = "투자 아이디어"
RISKS_HEADING = "투자 리스크"
TARGET_HEADING = "목표가"
NO_INFO = "자료 내 명시 없음"
SOURCE_RE = re.compile(r"\[(?:출처|source)\s*:\s*([^\]]+)\]", re.IGNORECASE)
SUMMARY_SECTION_KEYS = ("overview", "thesis", "risks")
DEFAULT_MAX_ITEMS = {"overview": 3, "thesis": 8, "risks": 8}
DEFAULT_COMPACT_RETRY_MAX_ITEMS = {"overview": 3, "thesis": 6, "risks": 6}

OCR_PROMPT = (
    "이 자료의 모든 텍스트를 그대로 추출하세요. 표, 수치, 축 라벨, 각주, 섹션 제목을 포함합니다.\n"
    "규칙:\n"
    "- 원문에 있는 문자와 숫자만 옮겨 적으세요. 요약·번역·해석·계산을 하지 마세요.\n"
    "- 목표주가·적정주가·밸류에이션 표가 있으면 항목명과 금액을 같은 줄에 붙여서 적으세요.\n"
    "  예: 'Base 목표주가 61,000원'\n"
    "- 읽을 수 없는 부분은 비워 두고 추측해서 채우지 마세요."
)
MAX_OCR_IMAGE_BYTES = 7 * 1024 * 1024
DEFAULT_MAX_OCR_SLIDES = 40
DEFAULT_MAX_OCR_IMAGES_PER_SLIDE = 4
DEFAULT_OCR_SLIDE_TEXT_THRESHOLD = 80
SLIDE_MARKER_RE = re.compile(r"^---\s*Slide\s+(\d+)\s*---\s*$", re.IGNORECASE)

SUMMARY_PROMPT = """You are an assistant analyst for a Korean investment study.
Use only the provided source blocks. Do not infer facts that are not in the materials.

Write in Korean unless the source itself is English-only.
Prefer concrete facts: numbers, dates, quantities, products, customers, contracts,
capacity, revenue, margins, valuation, and explicit schedules.

Return JSON that matches the provided response schema. Do not output markdown fences.

Required sections:
- overview: 기업 개요
- thesis: 투자 아이디어
- risks: 투자 리스크
- target_price: 목표가

Rules:
- overview must contain exactly these three slots:
  core_bm, current_market_position, future_growth_momentum.
- core_bm renders as "핵심 BM".
- current_market_position renders as "현재 시장 지위".
- future_growth_momentum renders as "앞으로의 성장 모멘텀".
- Each overview slot must contain 1 to 3 Korean sentences when supported.
- If an overview slot is not supported, set fact to "자료 내 명시 없음" and source to empty.
- 투자 아이디어 and 투자 리스크 are not fixed to 3 items. Include only supported items from the materials.
- If no supported investment idea or risk exists, write "자료 내 명시 없음" in that section.
- Every factual item must include a source value copied from the allowed source labels.
- Do not mention a source unless that exact source block actually supports the claim.
- The source field must be one of the allowed labels. Do not translate or reformat it.
- For "자료 내 명시 없음" items, leave source empty.
- Keep each fact to one concise sentence. Do not emit headings, labels-only items, or repeated explanations.
- Avoid duplicate claims across sections and within the same section.
- If the same claim is supported by multiple sources, prefer presentation materials:
  PPTX/PDF first, then presentation-script TXT, then voice-recording DOCX.
- Aim to stay within these section limits while keeping the most material facts:
{item_limits}

Company: {company}
Ticker: {ticker}
Presenter: {presenter}
Presentation month: {presentation_month}

Allowed source labels:
{allowed_labels}
"""

REPAIR_PROMPT = """The previous answer failed validation:
{errors}

Regenerate the summary from the source blocks as one complete compact JSON object using
the same response schema. Do not reuse partial or malformed JSON from the previous answer.
Keep only claims supported by the provided source blocks. Every factual item must use one
source copied from the allowed source labels below. Do not translate, shorten, or reformat
source labels. Keep each fact to one concise sentence.
Avoid duplicate claims and prefer presentation-material sources in this order:
PPTX/PDF first, then presentation-script TXT, then voice-recording DOCX.
Aim to stay within these section limits:
{item_limits}

Allowed source labels:
{allowed_labels}
"""

COMPACT_RETRY_PROMPT = """The previous answer was truncated before valid JSON completed:
{errors}

Regenerate the summary from the source blocks as one complete compact JSON object using
the same response schema. Keep only the most material non-duplicative claims. Use one
concise sentence per fact, omit lower-priority repeats, and never reuse partial JSON from
the previous answer.
If the same claim is supported by multiple sources, prefer presentation-material sources
in this order: PPTX/PDF first, then presentation-script TXT, then voice-recording DOCX.
Stay within these stricter section limits:
{item_limits}

Allowed source labels:
{allowed_labels}
"""


@dataclass
class SourceDocument:
    title: str
    content_block: dict[str, Any]
    block_labels: list[str]


@dataclass
class CitationText:
    text: str
    citations: list[Any]


def _strip_bullet_prefix(line: str) -> str:
    return re.sub(r"^\s*(?:[-*]\s+|\d+[\.)]\s+)", "", line).strip()


def _strip_markdown_emphasis(line: str) -> str:
    line = line.strip()
    line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
    line = re.sub(r"__(.+?)__", r"\1", line)
    return line.strip()


def _clean_summary_line(line: str) -> str:
    line = _strip_bullet_prefix(line)
    return _strip_markdown_emphasis(line).strip()


def _heading_key(line: str) -> str:
    line = line.strip()
    line = re.sub(r"^#+\s*", "", line)
    line = _clean_summary_line(line)
    line = line.rstrip(":：").strip()
    normalized = re.sub(r"\s+", "", line)
    if normalized in {"기업개요", "회사개요", "개요"} or normalized.startswith("기업개요"):
        return "overview"
    if "투자아이디어" in normalized or normalized in {"Upside", "투자포인트"}:
        return "thesis"
    if "투자리스크" in normalized or normalized in {"Downside", "리스크"}:
        return "risks"
    if "목표가" in normalized or "목표주가" in normalized or normalized in {"타겟프라이스", "TargetPrice"}:
        return "target_price"
    return ""


def _is_section_heading(line: str) -> bool:
    return bool(_heading_key(line))


def _has_source(line: str) -> bool:
    return bool(SOURCE_RE.search(line))


def _is_no_info(line: str) -> bool:
    return NO_INFO in line or "(empty)" in line


def _is_structural_line(line: str) -> bool:
    raw = line.strip()
    clean = _clean_summary_line(line).strip()
    normalized = re.sub(r"\s+", "", clean.rstrip(":："))
    if not clean or clean in {"---", "***", "___"}:
        return True
    if _is_section_heading(clean):
        return True
    if re.fullmatch(r"\*\*.+?[:：]\*\*", raw):
        return True
    if normalized in {
        "핵심BM",
        "시장지위",
        "성장모멘텀",
        "InvestmentIdea",
        "InvestmentRisk",
    }:
        return True
    if normalized in {"또한", "그리고", "다만", "반면", "한편"}:
        return True
    return False


def _citation_value(citation: Any, name: str, default: Any = None) -> Any:
    if isinstance(citation, dict):
        return citation.get(name, default)
    return getattr(citation, name, default)


def _citation_source(citation: Any, documents: list[SourceDocument]) -> str:
    citation_type = str(_citation_value(citation, "type", ""))
    doc_index = _citation_value(citation, "document_index", None)
    try:
        doc_index_int = int(doc_index)
    except (TypeError, ValueError):
        doc_index_int = -1

    title = str(_citation_value(citation, "document_title", "") or "")
    if 0 <= doc_index_int < len(documents):
        title = title or documents[doc_index_int].title

    if citation_type == "page_location":
        page = _citation_value(citation, "start_page_number", None)
        return f"{title}/p.{page}" if page is not None else title

    if citation_type == "content_block_location":
        block_index = _citation_value(citation, "start_block_index", None)
        try:
            block_index_int = int(block_index)
        except (TypeError, ValueError):
            block_index_int = -1
        if 0 <= doc_index_int < len(documents):
            labels = documents[doc_index_int].block_labels
            if 0 <= block_index_int < len(labels):
                return labels[block_index_int]
        return title

    return title


def _source_marker(citations: list[Any], documents: list[SourceDocument]) -> str:
    sources: list[str] = []
    for citation in citations:
        source = _citation_source(citation, documents).strip()
        if source and source not in sources:
            sources.append(source)
    if not sources:
        return ""
    return f"[출처: {'; '.join(sources)}]"


def _line_with_marker(line: str, marker: str) -> str:
    stripped = line.strip()
    if not stripped or _is_section_heading(stripped) or _has_source(stripped) or not marker:
        return stripped
    return f"{stripped} {marker}"


def _message_text_blocks(message: Any) -> list[CitationText]:
    blocks: list[CitationText] = []
    for block in getattr(message, "content", []) or []:
        text = getattr(block, "text", None)
        if text is None and isinstance(block, dict):
            text = block.get("text")
        if not text:
            continue
        citations = getattr(block, "citations", None)
        if citations is None and isinstance(block, dict):
            citations = block.get("citations")
        blocks.append(CitationText(text=str(text), citations=list(citations or [])))
    return blocks


def _annotated_response_text(message: Any, documents: list[SourceDocument]) -> str:
    """Compatibility helper for provider citation tests and saved debug data."""
    lines: list[str] = []
    for block in _message_text_blocks(message):
        marker = _source_marker(block.citations, documents)
        for raw in block.text.splitlines():
            line = raw.rstrip()
            if not line.strip():
                lines.append("")
            else:
                lines.append(_line_with_marker(line, marker))
    return "\n".join(lines).strip()


def parse_summary_markdown(
    text: str,
    company: str,
    ticker: str = "",
    presenter: str = "",
    presentation_month: str = "",
) -> Summary:
    """Parse legacy markdown summaries kept for compatibility and tests."""
    sections: dict[str, list[str]] = {
        "overview": [],
        "thesis": [],
        "risks": [],
        "target_price": [],
    }
    current = ""

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        key = _heading_key(line)
        if key:
            current = key
            continue
        if not current:
            continue
        if _is_structural_line(line):
            continue
        sections[current].append(_clean_summary_line(line))

    return Summary(
        company=company,
        ticker=ticker,
        presenter=presenter,
        presentation_month=presentation_month,
        overview="\n".join(sections["overview"]).strip(),
        thesis="\n".join(sections["thesis"]).strip(),
        risks="\n".join(sections["risks"]).strip(),
        conclusion="",
        target_price="\n".join(sections["target_price"]).strip(),
        raw_excerpt=text[:300],
    )


def _extract_source_values(line: str) -> list[str]:
    values: list[str] = []
    for match in SOURCE_RE.finditer(line):
        for source in re.split(r"\s*;\s*", match.group(1)):
            source = source.strip()
            if source:
                values.append(source)
    return values


def _canonical_source_label(value: str) -> str:
    value = str(value or "").strip()
    source_match = SOURCE_RE.search(value)
    if source_match:
        value = source_match.group(1)
    value = value.strip().strip("[](){}")
    value = value.replace("\\", "/")
    value = re.sub(r"\bslides?\b", "slide", value, flags=re.IGNORECASE)
    value = re.sub(r"\bpages?\b", "page", value, flags=re.IGNORECASE)
    value = re.sub(r"\bsections?\b", "section", value, flags=re.IGNORECASE)
    value = re.sub(r"\bparts?\b", "part", value, flags=re.IGNORECASE)
    value = re.sub(r"\b시트\b", "sheet", value, flags=re.IGNORECASE)
    value = value.replace("슬라이드", "slide")
    value = value.replace("페이지", "page")
    value = value.replace("섹션", "section")
    value = value.replace("파트", "part")
    value = re.sub(r"(slide|page|section|part)\s*0*([0-9]+)", r"\1\2", value, flags=re.IGNORECASE)
    value = re.sub(r"\s*/\s*", "/", value)
    return re.sub(r"[\s\[\]\(\){}:：_\-]+", "", value.casefold())


def _match_source_label(source: str, allowed_labels: list[str]) -> tuple[str, str]:
    source = str(source or "").strip()
    if not source:
        return "", "출처가 없습니다"

    canonical = _canonical_source_label(source)
    if not canonical:
        return "", f"출처가 비어 있습니다: {source}"

    canonical_map: dict[str, list[str]] = {}
    for label in allowed_labels:
        canonical_map.setdefault(_canonical_source_label(label), []).append(label)

    exact = canonical_map.get(canonical, [])
    if len(exact) == 1:
        return exact[0], ""
    if len(exact) > 1:
        return "", f"출처가 여러 라벨과 일치합니다: {source}"

    scored = sorted(
        (
            (SequenceMatcher(None, canonical, _canonical_source_label(label)).ratio(), label)
            for label in allowed_labels
        ),
        reverse=True,
    )
    if not scored or scored[0][0] < 0.86:
        return "", f"허용되지 않은 출처입니다: {source}"
    if len(scored) > 1 and scored[0][0] - scored[1][0] < 0.08:
        return "", f"출처가 애매합니다: {source}"
    return scored[0][1], ""


def validate_cited_summary(summary: Summary, allowed_labels: list[str] | None = None) -> list[str]:
    errors: list[str] = []
    if not (summary.overview or summary.thesis or summary.risks):
        errors.append("요약 본문이 비어 있습니다.")

    for section_name, body in [
        ("기업 개요", summary.overview),
        ("투자 아이디어", summary.thesis),
        ("투자 리스크", summary.risks),
    ]:
        for line in body.splitlines():
            line = line.strip()
            if not line or _is_no_info(line) or _is_structural_line(line):
                continue
            sources = _extract_source_values(line)
            if not sources:
                errors.append(f"{section_name} 항목에 출처가 없습니다: {line}")
                continue
            if allowed_labels is not None:
                for source in sources:
                    _matched, error = _match_source_label(source, allowed_labels)
                    if error:
                        errors.append(f"{section_name} 항목의 출처가 유효하지 않습니다: {error}")

    target = summary.target_price.strip()
    if target and not _is_no_info(target) and not _is_structural_line(target):
        sources = _extract_source_values(target)
        if not sources:
            errors.append(f"목표가 항목에 출처가 없습니다: {target}")
        elif allowed_labels is not None:
            for source in sources:
                _matched, error = _match_source_label(source, allowed_labels)
                if error:
                    errors.append(f"목표가 항목의 출처가 유효하지 않습니다: {error}")
    return errors


def _safe_debug_name(company: str, ticker: str) -> str:
    base = "_".join(part for part in [company, ticker] if part).strip() or "llm_summary"
    base = re.sub(r"[\\/:*?\"<>|\s]+", "_", base)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return f"{base}_{stamp}.json"


def _marker_label(path: Path, marker: str, part_index: int = 1) -> str:
    marker = marker.strip()
    match = re.match(r"^---\s*(Slide|Page|Section)\s*:\s*(.+?)\s*---$", marker, flags=re.IGNORECASE)
    if match:
        return f"{path.name}/{match.group(1).title()} {match.group(2)}"
    match = re.match(r"^---\s*(Slide|Page|Section)\s+(.+?)\s*---$", marker, flags=re.IGNORECASE)
    if match:
        return f"{path.name}/{match.group(1).title()} {match.group(2)}"
    match = re.match(r"^---\s*Sheet:\s*(.+?)\s*---$", marker, flags=re.IGNORECASE)
    if match:
        return f"{path.name}/Sheet {match.group(1)}"
    return f"{path.name}/part {part_index}"


def _split_text_chunks(path: Path, text: str, max_chars: int = 3500) -> tuple[list[dict[str, str]], list[str]]:
    chunks: list[dict[str, str]] = []
    labels: list[str] = []
    current_label = f"{path.name}/part 1"
    current_lines: list[str] = []
    part_index = 1

    def flush() -> None:
        nonlocal part_index, current_lines, current_label
        body = "\n".join(current_lines).strip()
        if not body:
            current_lines = []
            return
        while len(body) > max_chars:
            piece = body[:max_chars].rstrip()
            chunks.append({"label": current_label, "text": piece})
            labels.append(current_label)
            body = body[max_chars:].lstrip()
            part_index += 1
            current_label = f"{path.name}/part {part_index}"
        if body:
            chunks.append({"label": current_label, "text": body})
            labels.append(current_label)
        current_lines = []

    for raw in text.splitlines():
        line = raw.rstrip()
        if re.match(r"^---\s*(?:Slide|Page|Section|Sheet:)", line, flags=re.IGNORECASE):
            flush()
            part_index += 1
            current_label = _marker_label(path, line, part_index)
            continue
        current_lines.append(line)
    flush()

    if not chunks and text.strip():
        chunks.append({"label": path.name, "text": text.strip()})
        labels.append(path.name)
    return chunks, labels


def _read_text_document(path: Path, text: str | None = None) -> SourceDocument:
    body = text if text is not None else read_file(path)
    chunks, labels = _split_text_chunks(path, body)
    if not chunks:
        raise RuntimeError("no extractable text")
    return SourceDocument(
        title=path.name,
        block_labels=labels,
        content_block={
            "type": "source_blocks",
            "chunks": chunks,
        },
    )


def _native_text_by_slide(text: str) -> dict[int, str]:
    """`--- Slide N ---` 마커로 구분된 추출 텍스트를 슬라이드 번호별로 나눈다."""
    by_slide: dict[int, list[str]] = {}
    current: int | None = None
    for raw in (text or "").splitlines():
        marker = SLIDE_MARKER_RE.match(raw.strip())
        if marker:
            current = int(marker.group(1))
            by_slide.setdefault(current, [])
            continue
        if current is not None:
            by_slide[current].append(raw)
    return {index: "\n".join(lines).strip() for index, lines in by_slide.items()}


def _pptx_slide_images(path: Path) -> list[tuple[int, list[tuple[bytes, str]]]]:
    """슬라이드별 그림 blob 목록. 그룹 도형 안쪽까지 재귀한다."""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:
        raise RuntimeError("python-pptx is not installed") from e

    prs = Presentation(str(path))
    slides: list[tuple[int, list[tuple[bytes, str]]]] = []
    for index, slide in enumerate(prs.slides, start=1):
        images: list[tuple[bytes, str]] = []

        def walk(shapes) -> None:
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                    continue
                try:
                    image = shape.image
                except (AttributeError, ValueError):
                    continue
                blob = getattr(image, "blob", None)
                if blob:
                    images.append((blob, getattr(image, "content_type", "image/png") or "image/png"))

        walk(slide.shapes)
        slides.append((index, images))
    return slides


def _read_pdf_block(path: Path) -> SourceDocument:
    """Read PDFs through the local extractor so source labels stay uniform."""
    return _read_text_document(path)


def _document_blocks_from_files(file_paths: list[Path]) -> list[SourceDocument]:
    documents: list[SourceDocument] = []
    failures: list[str] = []
    for path in file_paths:
        try:
            documents.append(_read_text_document(path))
        except Exception as e:
            failures.append(f"{path.name}: {e}")
            continue
    if not documents and failures:
        raise RuntimeError("no readable source files for LLM summarization: " + "; ".join(failures))
    return documents


def _allowed_labels(documents: list[SourceDocument]) -> list[str]:
    labels: list[str] = []
    for document in documents:
        for label in document.block_labels:
            if label not in labels:
                labels.append(label)
    return labels


def _format_allowed_labels(labels: list[str]) -> str:
    return "\n".join(f"- {label}" for label in labels) or "- (none)"


def _format_item_limits(limits: dict[str, int]) -> str:
    return "\n".join(
        [
            f"- overview: at most {limits['overview']} items",
            f"- thesis: at most {limits['thesis']} items",
            f"- risks: at most {limits['risks']} items",
        ]
    )


def _documents_prompt_text(documents: list[SourceDocument]) -> str:
    blocks: list[str] = []
    for document in documents:
        for chunk in document.content_block.get("chunks", []):
            label = str(chunk.get("label") or document.title)
            text = str(chunk.get("text") or "").strip()
            if not text:
                continue
            blocks.append(
                f"--- BLOCK START: [{label}] ---\n"
                f"{text}\n"
                f"--- BLOCK END: [{label}] ---"
            )
    return "\n\n".join(blocks)


def _summary_response_schema() -> dict[str, Any]:
    try:
        from pydantic import BaseModel, ConfigDict, Field

        class SummaryItemModel(BaseModel):
            model_config = ConfigDict(extra="forbid")

            fact: str = Field(description="Supported factual sentence or '자료 내 명시 없음'.")
            source: str = Field(
                description="One allowed source label copied exactly, or empty for '자료 내 명시 없음'."
            )

        class OverviewResponseModel(BaseModel):
            model_config = ConfigDict(extra="forbid")

            core_bm: SummaryItemModel
            current_market_position: SummaryItemModel
            future_growth_momentum: SummaryItemModel

        class SummaryResponseModel(BaseModel):
            model_config = ConfigDict(extra="forbid")

            overview: OverviewResponseModel
            thesis: list[SummaryItemModel]
            risks: list[SummaryItemModel]
            target_price: list[SummaryItemModel]

        return SummaryResponseModel.model_json_schema()
    except ImportError:
        pass

    item_schema = {
        "type": "object",
        "properties": {
            "fact": {
                "type": "string",
                "description": "Supported factual sentence or '자료 내 명시 없음'.",
            },
            "source": {
                "type": "string",
                "description": "One allowed source label copied exactly, or empty for '자료 내 명시 없음'.",
            },
        },
        "required": ["fact", "source"],
        "additionalProperties": False,
    }
    overview_schema = {
        "type": "object",
        "properties": {
            "core_bm": item_schema,
            "current_market_position": item_schema,
            "future_growth_momentum": item_schema,
        },
        "required": ["core_bm", "current_market_position", "future_growth_momentum"],
        "additionalProperties": False,
    }
    return {
        "type": "object",
        "properties": {
            "overview": overview_schema,
            "thesis": {"type": "array", "items": item_schema},
            "risks": {"type": "array", "items": item_schema},
            "target_price": {"type": "array", "items": item_schema},
        },
        "required": ["overview", "thesis", "risks", "target_price"],
        "additionalProperties": False,
    }


def _strip_code_fence(text: str) -> str:
    text = (text or "").strip()
    match = re.match(r"^```(?:json)?\s*(.+?)\s*```$", text, flags=re.DOTALL | re.IGNORECASE)
    return match.group(1).strip() if match else text


def _parse_response_json(text: str) -> dict[str, Any]:
    data = json.loads(_strip_code_fence(text))
    if not isinstance(data, dict):
        raise ValueError("Gemini response JSON is not an object")
    return data


def _json_error_context(text: str, error: json.JSONDecodeError, radius: int = 80) -> str:
    start = max(0, error.pos - radius)
    end = min(len(text), error.pos + radius)
    snippet = text[start:end].replace("\n", "\\n")
    return f"{error.msg} at line {error.lineno} column {error.colno}; near={snippet!r}"


def _response_metadata(response: Any) -> dict[str, Any]:
    candidates = getattr(response, "candidates", None) or []
    first_candidate = candidates[0] if candidates else None
    finish_reason = getattr(first_candidate, "finish_reason", None) if first_candidate is not None else None
    if hasattr(finish_reason, "value"):
        finish_reason = finish_reason.value
    elif finish_reason is not None:
        finish_reason = str(finish_reason)

    usage = getattr(response, "usage_metadata", None)
    usage_payload: dict[str, Any] = {}
    for name in [
        "prompt_token_count",
        "candidates_token_count",
        "thoughts_token_count",
        "total_token_count",
        "cached_content_token_count",
    ]:
        value = getattr(usage, name, None) if usage is not None else None
        if value is not None:
            usage_payload[name] = value

    return {
        "finish_reason": finish_reason,
        "usage_metadata": usage_payload,
    }


def _coerce_items(payload: dict[str, Any], key: str) -> list[dict[str, str]]:
    raw = payload.get(key, [])
    if isinstance(raw, dict):
        raw = [raw]
    if isinstance(raw, str):
        raw = [{"fact": raw, "source": ""}]
    if not isinstance(raw, list):
        return []

    items: list[dict[str, str]] = []
    for item in raw:
        if isinstance(item, str):
            items.append({"fact": item, "source": ""})
            continue
        if not isinstance(item, dict):
            continue
        source = item.get("source", "")
        if isinstance(source, list):
            source = "; ".join(str(part) for part in source)
        items.append(
            {
                "fact": str(item.get("fact", "")).strip(),
                "source": str(source or "").strip(),
            }
        )
    return items


def _coerce_overview_items(payload: dict[str, Any]) -> list[dict[str, str]]:
    raw = payload.get("overview", {})
    if isinstance(raw, list):
        return _coerce_items(payload, "overview")
    if not isinstance(raw, dict):
        return []
    def item(value: Any) -> dict[str, str]:
        if isinstance(value, dict):
            return {
                "fact": str(value.get("fact", "")).strip(),
                "source": str(value.get("source", "") or "").strip(),
            }
        if isinstance(value, str):
            return {"fact": value.strip(), "source": ""}
        return {"fact": "", "source": ""}
    return [
        item(raw.get("core_bm")),
        item(raw.get("current_market_position")),
        item(raw.get("future_growth_momentum")),
    ]


def _sentence_count(text: str) -> int:
    clean = re.sub(r"\[[^\]]+\]", "", text or "").strip()
    if not clean or _is_no_info(clean):
        return 0
    parts = [part for part in re.split(r"(?<=[.!?])\s+", clean) if part.strip()]
    return len(parts) if parts else 1


def _overview_lines_from_items(
    items: list[dict[str, str]],
    allowed_labels: list[str],
) -> tuple[list[str], list[str]]:
    labels = ["핵심 BM", "현재 시장 지위", "앞으로의 성장 모멘텀"]
    padded = items[:3] + [{"fact": NO_INFO, "source": ""}] * max(0, 3 - len(items))
    lines: list[str] = []
    errors: list[str] = []
    for label, item in zip(labels, padded):
        fact = _clean_summary_line(item.get("fact", "")) or NO_INFO
        source = item.get("source", "").strip()
        if _is_no_info(fact):
            lines.append(f"{label}: {NO_INFO}")
            continue
        sentence_count = _sentence_count(fact)
        if sentence_count > 3:
            errors.append(f"기업 개요 항목이 3문장을 초과합니다: {label}")
        matched, error = _match_source_label(source, allowed_labels)
        if error:
            errors.append(f"기업 개요 항목의 출처가 유효하지 않습니다: {error} / label={label}")
            lines.append(f"{label}: {fact}")
            continue
        lines.append(f"{label}: {fact} [출처: {matched}]")
    return lines, errors


def _normalized_fact_key(fact: str) -> str:
    fact = _clean_summary_line(fact)
    fact = re.sub(r"\[[^\]]+\]", "", fact)
    fact = re.sub(r"[^\w가-힣]+", "", fact.casefold())
    return fact


def _facts_are_duplicates(left: str, right: str) -> bool:
    left_key = _normalized_fact_key(left)
    right_key = _normalized_fact_key(right)
    if not left_key or not right_key:
        return False
    if left_key == right_key:
        return True
    return SequenceMatcher(None, left_key, right_key).ratio() >= 0.9


def _source_priority(source: str) -> int:
    lower = str(source or "").casefold()
    if ".pptx/" in lower or ".pdf/" in lower:
        return 0
    if "발표대본" in lower or ".txt/" in lower:
        return 1
    if "음성녹음" in lower or ".docx/" in lower:
        return 2
    return 3


def _prefer_item(left: dict[str, str], right: dict[str, str]) -> dict[str, str]:
    left_rank = _source_priority(left.get("source", ""))
    right_rank = _source_priority(right.get("source", ""))
    if right_rank < left_rank:
        return right
    if left_rank < right_rank:
        return left
    if len(right.get("fact", "")) > len(left.get("fact", "")):
        return right
    return left


def _dedupe_items(items: list[dict[str, str]]) -> list[dict[str, str]]:
    deduped: list[dict[str, str]] = []
    for item in items:
        fact = item.get("fact", "").strip()
        if not fact:
            continue
        if _is_no_info(fact):
            if not deduped:
                deduped.append(item)
            continue
        duplicate_index = next(
            (
                index
                for index, existing in enumerate(deduped)
                if _facts_are_duplicates(existing.get("fact", ""), fact)
            ),
            None,
        )
        if duplicate_index is None:
            deduped.append(item)
            continue
        deduped[duplicate_index] = _prefer_item(deduped[duplicate_index], item)
    return deduped


def _cap_items(items: list[dict[str, str]], limit: int) -> list[dict[str, str]]:
    if limit <= 0:
        return []
    factual_items = [item for item in items if not _is_no_info(item.get("fact", ""))]
    if factual_items:
        return factual_items[:limit]
    return items[:1]


def _prepare_structured_payload(
    payload: dict[str, Any],
    max_items: dict[str, int],
) -> tuple[dict[str, Any], dict[str, dict[str, int]]]:
    prepared = dict(payload)
    stats: dict[str, dict[str, int]] = {}
    for key in SUMMARY_SECTION_KEYS:
        items = _coerce_overview_items(payload) if key == "overview" else _coerce_items(payload, key)
        if key == "overview":
            deduped = items
            capped = items[:3]
            prepared[key] = {
                "core_bm": capped[0] if len(capped) >= 1 else {"fact": NO_INFO, "source": ""},
                "current_market_position": capped[1] if len(capped) >= 2 else {"fact": NO_INFO, "source": ""},
                "future_growth_momentum": capped[2] if len(capped) >= 3 else {"fact": NO_INFO, "source": ""},
            }
        else:
            deduped = _dedupe_items(items)
            capped = _cap_items(deduped, max_items[key])
            prepared[key] = capped
        stats[key] = {
            "before_dedupe": len(items),
            "after_dedupe": len(deduped),
            "after_cap": len(capped),
        }
    prepared["target_price"] = _coerce_items(payload, "target_price")
    return prepared, stats


def _lines_from_structured_items(
    section_name: str,
    items: list[dict[str, str]],
    allowed_labels: list[str],
) -> tuple[list[str], list[str]]:
    lines: list[str] = []
    errors: list[str] = []
    for item in items:
        fact = _clean_summary_line(item.get("fact", ""))
        source = item.get("source", "").strip()
        if not fact:
            continue
        if _is_no_info(fact):
            lines.append(fact)
            continue
        matched, error = _match_source_label(source, allowed_labels)
        if error:
            errors.append(f"{section_name} 항목의 출처가 유효하지 않습니다: {error} / fact={fact}")
            lines.append(fact)
            continue
        lines.append(f"{fact} [출처: {matched}]")
    return lines, errors


def structured_payload_to_summary(
    payload: dict[str, Any],
    allowed_labels: list[str],
    company: str,
    ticker: str = "",
    presenter: str = "",
    presentation_month: str = "",
    raw_excerpt: str = "",
) -> tuple[Summary, list[str]]:
    overview, overview_errors = _overview_lines_from_items(
        _coerce_overview_items(payload),
        allowed_labels,
    )
    thesis, thesis_errors = _lines_from_structured_items(
        "투자 아이디어", _coerce_items(payload, "thesis"), allowed_labels
    )
    risks, risks_errors = _lines_from_structured_items(
        "투자 리스크", _coerce_items(payload, "risks"), allowed_labels
    )
    target_price, target_errors = _lines_from_structured_items(
        "목표가", _coerce_items(payload, "target_price"), allowed_labels
    )

    summary = Summary(
        company=company,
        ticker=ticker,
        presenter=presenter,
        presentation_month=presentation_month,
        overview="\n".join(overview).strip(),
        thesis="\n".join(thesis).strip(),
        risks="\n".join(risks).strip(),
        conclusion="",
        target_price="\n".join(target_price).strip(),
        raw_excerpt=raw_excerpt[:300],
    )
    errors = overview_errors + thesis_errors + risks_errors + target_errors
    errors.extend(validate_cited_summary(summary, allowed_labels=allowed_labels))
    return summary, errors


def _is_retryable_api_error(error: Exception) -> bool:
    status = getattr(error, "status_code", None) or getattr(error, "code", None)
    try:
        status_int = int(status)
    except (TypeError, ValueError):
        status_int = 0
    if status_int == 429 or 500 <= status_int <= 599:
        return True
    text = str(error).lower()
    return any(
        marker in text
        for marker in [
            "429",
            "too many requests",
            "rate limit",
            "quota",
            "500",
            "502",
            "503",
            "504",
            "timeout",
            "temporarily unavailable",
            "connection",
        ]
    )


def _parse_retry_delay(error: Exception) -> float | None:
    """Gemini 에러 응답에서 retryDelay 값(초)을 파싱한다."""
    import re
    text = str(error)
    match = re.search(r"retrydelay['\"]:\s*['\"](\d+\.?\d*)s", text, re.IGNORECASE)
    if match:
        return float(match.group(1))
    return None


def _normalized_item_limits(
    configured: dict[str, int] | None,
    defaults: dict[str, int],
) -> dict[str, int]:
    limits = dict(defaults)
    if not configured:
        return limits
    for key in SUMMARY_SECTION_KEYS:
        value = configured.get(key)
        if value is None:
            continue
        limits[key] = max(0, int(value))
    return limits


class LLMSummarizer(Summarizer):
    """Gemini API summarizer with structured source-label validation."""

    def __init__(
        self,
        model: str = "gemini-2.5-flash",
        api_key: str | None = None,
        max_tokens: int = 3000,
        max_retries: int = 1,
        max_items: dict[str, int] | None = None,
        compact_retry_max_items: dict[str, int] | None = None,
        debug_dir: str | Path | None = None,
        fallback_on_error: bool = False,
        api_retries: int = 3,
        sleep_func: Callable[[float], None] = time.sleep,
        max_ocr_slides: int = DEFAULT_MAX_OCR_SLIDES,
        max_ocr_images_per_slide: int = DEFAULT_MAX_OCR_IMAGES_PER_SLIDE,
        ocr_slide_text_threshold: int = DEFAULT_OCR_SLIDE_TEXT_THRESHOLD,
    ) -> None:
        self.model = model
        self.api_key = api_key or os.environ.get("GEMINI_API_KEY")
        self.max_tokens = max_tokens
        self.max_retries = max_retries
        self.max_items = _normalized_item_limits(max_items, DEFAULT_MAX_ITEMS)
        self.compact_retry_max_items = _normalized_item_limits(
            compact_retry_max_items,
            DEFAULT_COMPACT_RETRY_MAX_ITEMS,
        )
        self.debug_dir = Path(debug_dir) if debug_dir else None
        self.fallback_on_error = fallback_on_error
        self.api_retries = api_retries
        self.sleep_func = sleep_func
        self.max_ocr_slides = max_ocr_slides
        self.max_ocr_images_per_slide = max_ocr_images_per_slide
        self.ocr_slide_text_threshold = ocr_slide_text_threshold
        self._genai_client: Any | None = None
        self._last_response_metadata: dict[str, Any] = {}
        self._last_attempt_mode = "standard"
        self._last_item_stats: dict[str, dict[str, int]] = {}

    def _client(self) -> Any:
        try:
            from google import genai
        except ImportError as e:
            raise RuntimeError("google-genai is not installed; run pip install google-genai") from e

        if not self.api_key:
            raise RuntimeError("GEMINI_API_KEY is not set")
        if self._genai_client is None:
            self._genai_client = genai.Client(api_key=self.api_key)
        return self._genai_client

    def _read_pdf_via_gemini(self, path: Path) -> SourceDocument:
        """pypdf로 텍스트 추출 불가한 스캔 PDF를 Gemini Files API로 OCR 처리."""
        print(f"  [ocr] Gemini Vision으로 PDF 텍스트 추출 중: {path.name}")
        client = self._client()
        with open(path, "rb") as f:
            file_ref = client.files.upload(file=f, config={"mime_type": "application/pdf", "display_name": "source.pdf"})
        response = client.models.generate_content(
            model=self.model,
            contents=[
                file_ref,
                OCR_PROMPT,
            ],
        )
        text = str(getattr(response, "text", "") or "")
        chunks, labels = _split_text_chunks(path, text)
        if not chunks:
            raise RuntimeError(f"Gemini OCR가 텍스트를 반환하지 않음: {path.name}")
        return SourceDocument(
            title=path.name,
            block_labels=labels,
            content_block={"type": "source_blocks", "chunks": chunks},
        )

    def _read_pptx_via_gemini(self, path: Path, native_text: str = "") -> SourceDocument:
        """이미지 위주 PPTX를 슬라이드 단위로 OCR 한다.

        python-pptx가 뽑아낸 네이티브 텍스트(제목 등)는 그대로 두고,
        텍스트가 거의 없는 슬라이드의 그림만 Gemini Vision으로 읽어 합친다.
        슬라이드 라벨을 유지해야 출처 표기가 기존과 같은 형태로 남는다.
        """
        slides = _pptx_slide_images(path)
        if not slides:
            raise RuntimeError(f"OCR 대상 이미지를 찾지 못함: {path.name}")

        native_by_slide = _native_text_by_slide(native_text)
        targets = [
            (index, images)
            for index, images in slides
            if images and len(native_by_slide.get(index, "")) < self.ocr_slide_text_threshold
        ]
        if not targets:
            raise RuntimeError(f"OCR가 필요한 슬라이드가 없음: {path.name}")

        targets = targets[: self.max_ocr_slides]
        print(f"  [ocr] Gemini Vision으로 PPTX 슬라이드 {len(targets)}장 추출 중: {path.name}")

        client = self._client()
        from google.genai import types as genai_types

        ocr_by_slide: dict[int, str] = {}
        for slide_index, images in targets:
            parts: list[Any] = []
            for blob, content_type in images[: self.max_ocr_images_per_slide]:
                if len(blob) > MAX_OCR_IMAGE_BYTES:
                    continue
                parts.append(genai_types.Part.from_bytes(data=blob, mime_type=content_type))
            if not parts:
                continue
            try:
                response = client.models.generate_content(
                    model=self.model,
                    contents=[*parts, OCR_PROMPT],
                )
            except Exception as e:  # noqa: BLE001 - 슬라이드 하나 실패로 전체를 버리지 않는다
                print(f"  [ocr] slide {slide_index} 실패: {e}", file=sys.stderr)
                continue
            text = str(getattr(response, "text", "") or "").strip()
            if text:
                ocr_by_slide[slide_index] = text

        if not ocr_by_slide:
            raise RuntimeError(f"Gemini OCR가 텍스트를 반환하지 않음: {path.name}")

        merged: list[str] = []
        for index, _images in slides:
            body = "\n".join(
                part for part in (native_by_slide.get(index, "").strip(), ocr_by_slide.get(index, "")) if part
            )
            if body:
                merged.append(f"--- Slide {index} ---\n{body}")

        chunks, labels = _split_text_chunks(path, "\n\n".join(merged))
        if not chunks:
            raise RuntimeError(f"Gemini OCR 결과가 비어 있음: {path.name}")
        return SourceDocument(
            title=path.name,
            block_labels=labels,
            content_block={"type": "source_blocks", "chunks": chunks},
        )

    def _read_document_with_ocr_fallback(self, path: Path) -> tuple[SourceDocument, str]:
        """일반 추출을 먼저 시도하고, 비었거나 희박하면 OCR로 보강한다.

        반환값의 두 번째 항목은 로그용 상태 문자열이다.
        """
        native_text = ""
        try:
            native_text = read_file(path)
        except Exception as e:  # noqa: BLE001
            native_text = ""
            print(f"  [warning] {path.name} 텍스트 추출 실패: {e}", file=sys.stderr)

        sparse = extraction_is_sparse(path, native_text)
        if not sparse:
            return _read_text_document(path, native_text), "text"

        suffix = path.suffix.lower()
        chars = len(native_text)
        print(f"  [ocr] {path.name} 텍스트가 희박함({chars:,}자) — OCR 시도")
        try:
            if suffix == ".pdf":
                return self._read_pdf_via_gemini(path), "ocr"
            if suffix == ".pptx":
                return self._read_pptx_via_gemini(path, native_text), "ocr"
        except Exception as e:  # noqa: BLE001
            print(f"  [warning] {path.name} OCR 실패: {e}", file=sys.stderr)

        # OCR이 안 되면 원래 추출 결과라도 쓴다. 그것도 비면 호출부에서 실패 처리.
        return _read_text_document(path, native_text), "text_sparse"

    def _call_api(self, documents: list[SourceDocument], prompt: str) -> str:
        contents = [
            _documents_prompt_text(documents),
            prompt,
        ]
        schema = _summary_response_schema()
        last_error: Exception | None = None
        for attempt in range(self.api_retries + 1):
            try:
                response = self._client().models.generate_content(
                    model=self.model,
                    contents=contents,
                    config={
                        "max_output_tokens": self.max_tokens,
                        "response_mime_type": "application/json",
                        "response_json_schema": schema,
                    },
                )
                self._last_response_metadata = _response_metadata(response)
                return str(getattr(response, "text", "") or "")
            except Exception as e:
                last_error = e
                if attempt >= self.api_retries or not _is_retryable_api_error(e):
                    raise
                parsed = _parse_retry_delay(last_error)
                if parsed is not None:
                    delay = parsed + random.uniform(0, 1.0)
                elif _is_retryable_api_error(last_error):
                    delay = max(2**attempt * 5, 12.0) + random.uniform(0, 1.0)
                else:
                    delay = (2**attempt) + random.uniform(0, 0.25)
                self.sleep_func(delay)
        raise RuntimeError(f"Gemini API call failed: {last_error}")

    def _write_debug(
        self,
        company: str,
        ticker: str,
        response_text: str,
        errors: list[str],
        documents: list[SourceDocument],
        structured_payload: dict[str, Any] | None,
        allowed_labels: list[str],
    ) -> None:
        if self.debug_dir is None:
            return
        self.debug_dir.mkdir(parents=True, exist_ok=True)
        payload = {
            "company": company,
            "ticker": ticker,
            "model": self.model,
            "errors": errors,
            "response_text": response_text,
            "structured_payload": structured_payload,
            "response_metadata": self._last_response_metadata,
            "attempt_mode": self._last_attempt_mode,
            "item_stats": self._last_item_stats,
            "allowed_labels": allowed_labels,
            "documents": [
                {"title": doc.title, "block_labels": doc.block_labels}
                for doc in documents
            ],
        }
        path = self.debug_dir / _safe_debug_name(company, ticker)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    def _summarize_documents(
        self,
        documents: list[SourceDocument],
        company: str,
        ticker: str = "",
        presenter: str = "",
        presentation_month: str = "",
    ) -> Summary:
        allowed_labels = _allowed_labels(documents)
        label_text = _format_allowed_labels(allowed_labels)
        standard_limits_text = _format_item_limits(self.max_items)
        compact_limits_text = _format_item_limits(self.compact_retry_max_items)
        prompt = SUMMARY_PROMPT.format(
            company=company,
            ticker=ticker or "unknown ticker",
            presenter=presenter or "-",
            presentation_month=presentation_month or "-",
            allowed_labels=label_text,
            item_limits=standard_limits_text,
        )

        response_text = ""
        structured_payload: dict[str, Any] | None = None
        errors: list[str] = []
        raw_excerpt = _documents_prompt_text(documents)[:300]
        compact_retry_pending = False
        self._last_item_stats = {}
        for attempt in range(self.max_retries + 1):
            if attempt == 0:
                self._last_attempt_mode = "standard"
                active_limits = self.max_items
                call_prompt = prompt
            elif compact_retry_pending:
                self._last_attempt_mode = "compact_retry"
                active_limits = self.compact_retry_max_items
                call_prompt = COMPACT_RETRY_PROMPT.format(
                    errors="\n".join(errors),
                    allowed_labels=label_text,
                    item_limits=compact_limits_text,
                )
            else:
                self._last_attempt_mode = "repair"
                active_limits = self.max_items
                call_prompt = REPAIR_PROMPT.format(
                    errors="\n".join(errors),
                    allowed_labels=label_text,
                    item_limits=standard_limits_text,
                )
            try:
                response_text = self._call_api(documents, call_prompt)
                structured_payload = _parse_response_json(response_text)
            except json.JSONDecodeError as e:
                finish_reason = self._last_response_metadata.get("finish_reason")
                prefix = (
                    "Gemini response appears truncated at max_output_tokens"
                    if finish_reason == "MAX_TOKENS"
                    else "Gemini response JSON parsing failed"
                )
                errors = [f"{prefix}: {_json_error_context(response_text, e)}"]
                compact_retry_pending = finish_reason == "MAX_TOKENS"
                self._write_debug(
                    company,
                    ticker,
                    response_text,
                    errors,
                    documents,
                    structured_payload,
                    allowed_labels,
                )
                continue
            except ValueError as e:
                errors = [f"Gemini response JSON parsing failed: {e}"]
                compact_retry_pending = False
                self._write_debug(
                    company,
                    ticker,
                    response_text,
                    errors,
                    documents,
                    structured_payload,
                    allowed_labels,
                )
                continue

            structured_payload, self._last_item_stats = _prepare_structured_payload(
                structured_payload,
                active_limits,
            )
            summary, errors = structured_payload_to_summary(
                structured_payload,
                allowed_labels=allowed_labels,
                company=company,
                ticker=ticker,
                presenter=presenter,
                presentation_month=presentation_month,
                raw_excerpt=raw_excerpt,
            )
            if not errors:
                return summary
            compact_retry_pending = False

        self._write_debug(
            company,
            ticker,
            response_text,
            errors,
            documents,
            structured_payload,
            allowed_labels,
        )
        prefix = (
            "LLM summary failed after compact retry"
            if self._last_attempt_mode == "compact_retry"
            else "LLM summary failed citation validation"
        )
        raise RuntimeError(prefix + ": " + "; ".join(errors))

    def summarize_files(
        self,
        file_paths: list[str | Path],
        company: str,
        ticker: str = "",
        presenter: str = "",
        presentation_month: str = "",
    ) -> Summary:
        paths = [Path(path).resolve() for path in file_paths if Path(path).exists()]
        if not paths:
            raise RuntimeError("no existing source files for LLM summarization")

        documents: list[SourceDocument] = []
        failures: list[str] = []
        for path in paths:
            try:
                document, mode = self._read_document_with_ocr_fallback(path)
            except RuntimeError as e:
                failures.append(f"{path.name}: {e}")
                continue
            if mode == "text_sparse":
                print(
                    f"  [warning] {path.name} 은 이미지 위주라 텍스트가 거의 없고 OCR도 실패했습니다. "
                    "요약 품질이 낮을 수 있습니다.",
                    file=sys.stderr,
                )
            documents.append(document)
        if not documents:
            raise RuntimeError("no readable source files for LLM summarization: " + "; ".join(failures))

        return self._summarize_documents(documents, company, ticker, presenter, presentation_month)

    def summarize(
        self,
        text: str,
        company: str,
        ticker: str = "",
        presenter: str = "",
        presentation_month: str = "",
    ) -> Summary:
        if not text or not text.strip():
            return Summary(company=company, ticker=ticker, presenter=presenter, presentation_month=presentation_month)
        document = _read_text_document(Path("source.txt"), text)
        return self._summarize_documents([document], company, ticker, presenter, presentation_month)
